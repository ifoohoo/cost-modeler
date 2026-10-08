# -*- coding: utf-8 -*-
"""读取工作簿、上一版表 A、留存副本和机器导出名录。"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .model import (
    DASH, EXPORT_FILES, HEADERS, LOGICAL_SHEET, SEPARATOR_PREFIX, TABLE_A_GROUP, TABLE_A_SHEET,
    Options, Rec, Table, TEMPLATE_SHEETS,
)
from .textutil import unkey


def _text(value, number_format: str | None = None) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        fmt = number_format or ""
        if "0.0000" in fmt:
            return f"{value:.4f}"
        if "0.00" in fmt:
            return f"{value:.2f}"
        if abs(value - round(value)) < 1e-9:
            return str(int(round(value)))
        return format(value, "f").rstrip("0").rstrip(".")
    return str(value)


def _is_separator(values: list[str]) -> bool:
    return any(item.startswith(SEPARATOR_PREFIX) for item in values if item)


def read_grid(ws, *, skip_note: bool, require_separator: bool) -> tuple[list[str], list[tuple[int, list[str]]], bool]:
    """返回表头、数据行（Excel 行号 + 单元格文本）、是否缺分隔行。"""
    width = ws.max_column or 0
    headers = [_text(ws.cell(2, col).value) for col in range(1, width + 1)]
    while headers and headers[-1] == "":
        headers.pop()
    width = len(headers)
    separator_at = _find_separator(ws, width)
    missing = require_separator and separator_at is None
    rows = []
    if separator_at is None and require_separator:
        return headers, rows, True
    last = separator_at if separator_at else (ws.max_row or 0) + 1
    for row in range(3, last):
        if skip_note and row == 3:
            continue
        values = [_text(ws.cell(row, col).value, ws.cell(row, col).number_format) for col in range(1, width + 1)]
        if not any(values):
            continue
        if _is_separator(values):
            break
        rows.append((row, values))
    return headers, rows, missing


def _find_separator(ws, width: int) -> int | None:
    for row in range(1, (ws.max_row or 0) + 1):
        values = [_text(ws.cell(row, col).value) for col in range(1, width + 1)]
        if _is_separator(values):
            return row
    return None


def _role_columns(ws, width: int) -> tuple[list[int], bool]:
    """分组行里文字以「角色分工」开头的合并区域，其覆盖的列即角色列。"""
    for merged in sorted(ws.merged_cells.ranges, key=lambda item: (item.min_row, item.min_col)):
        if merged.min_row <= 2 <= merged.max_row:
            text = _text(ws.cell(merged.min_row, merged.min_col).value)
            if text.startswith(TABLE_A_GROUP):
                return list(range(merged.min_col, min(merged.max_col, width) + 1)), True
    return [], False


def read_table_a_grid(ws, *, require_separator: bool) -> tuple[list[str], list[tuple[int, list[str]]], bool, list[int], bool]:
    """表 A 双行表头：第 2 行分组、第 3 行列名、第 4 行填写说明、第 5 行起数据。"""
    width = ws.max_column or 0
    headers = [_text(ws.cell(3, col).value) for col in range(1, width + 1)]
    while headers and headers[-1] == "":
        headers.pop()
    width = len(headers)
    role_cols, group_found = _role_columns(ws, width)
    separator_at = _find_separator(ws, width)
    missing = require_separator and separator_at is None
    rows = []
    if separator_at is None and require_separator:
        return headers, rows, True, role_cols, group_found
    last = separator_at if separator_at else (ws.max_row or 0) + 1
    for row in range(5, last):
        values = [_text(ws.cell(row, col).value, ws.cell(row, col).number_format) for col in range(1, width + 1)]
        if not any(values):
            continue
        if _is_separator(values):
            break
        rows.append((row, values))
    return headers, rows, missing, role_cols, group_found


def _records(sheet_name: str, headers: list[str], raw_rows, case: str | None = None) -> list[Rec]:
    records = []
    for row_number, values in raw_rows:
        data = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
        records.append(Rec(data, row_number, sheet_name, case=case))
    return records


def _sheet_case(name: str, prefix: str) -> str | None:
    marker = prefix + "·"
    if name.startswith(marker):
        return name[len(marker):]
    return None


class WorkbookData:
    def __init__(self):
        self.tables: dict[str, Table] = {}
        self.by_logical: dict[str, list[Rec]] = {}
        self.catalog: set[str] | None = None
        self.catalog_missing_files: list[str] = []
        self.previous = None  # None 表示未提供；[] 表示提供了但没有行
        self.previous_headers: list[str] = []
        self.previous_role_cols: list[int] = []
        self.previous_group_found: bool = True
        self.copy_rows = None  # None 表示未提供

    def logical(self, key: str) -> list[Rec]:
        return self.by_logical.get(key, [])

    def readable(self, sink, code: str, *keys: str) -> bool:
        """查当前领域输入是否可读；分表正文缺失不能代表总量为零。"""
        from .messages import line, say
        blocked = []
        for key in keys:
            name = LOGICAL_SHEET[key]
            names = [name]
            if key in ("分摊", "恒等式"):
                names += [item for item in self.tables if item.startswith(name + "·")]
            for item in names:
                table = self.tables.get(item)
                # 未要求的可选表缺席由调用方按业务条件处理。
                if table is not None and table.blocked:
                    blocked.append(item)
        if blocked:
            text = line(code, say("check.blocked", p0="、".join(dict.fromkeys(blocked))))
            if text not in sink.unable.get(code, []):
                sink.set_unable(code, text)
        return not blocked


def _expected_headers(sheet_name: str) -> list[str] | None:
    if sheet_name in HEADERS:
        return HEADERS[sheet_name]
    for prefix in ("分摊表", "成本恒等式"):
        if sheet_name.startswith(prefix + "·"):
            return HEADERS[prefix]
    if sheet_name == "全系统成本恒等式":
        return HEADERS["成本恒等式"]
    if sheet_name == "已结期次留存副本":
        return HEADERS["分摊表"]
    return None


def load_workbook_file(path: Path) -> WorkbookData:
    book = WorkbookData()
    wb = load_workbook(path, data_only=True)
    present = list(wb.sheetnames)
    for name in TEMPLATE_SHEETS:
        if name == TABLE_A_SHEET:
            if name not in present:
                book.tables[name] = Table(name, [], present=False, separator_missing=True)
                continue
            headers, raw, missing, role_cols, group_found = read_table_a_grid(wb[name], require_separator=True)
            table = Table(name, headers, separator_missing=missing, role_cols=role_cols, role_group_found=group_found)
            if not table.blocked:
                table.rows = _records(name, headers, raw)
            book.tables[name] = table
            continue
        expected = HEADERS[name]
        if name not in present:
            book.tables[name] = Table(name, expected, present=False, separator_missing=True)
            continue
        headers, raw, missing = read_grid(wb[name], skip_note=True, require_separator=True)
        table = Table(name, headers, separator_missing=missing, header_bad=headers != expected)
        if not table.blocked:
            case = None
            table.rows = _records(name, headers, raw, case)
        book.tables[name] = table

    for name in present:
        if name in book.tables or name in ("说明", "整页级份额（机器执行）"):
            continue
        expected = _expected_headers(name)
        if expected is None:
            continue
        headers, raw, missing = read_grid(wb[name], skip_note=True, require_separator=True)
        table = Table(name, headers, separator_missing=missing, header_bad=headers != expected)
        if not table.blocked:
            case = _sheet_case(name, "分摊表") or _sheet_case(name, "成本恒等式")
            if name == "全系统成本恒等式":
                case = "全系统"
            table.rows = _records(name, headers, raw, case)
        book.tables[name] = table

    wb.close()
    _index(book)
    return book


def _index(book: WorkbookData) -> None:
    def take(name: str) -> list[Rec]:
        table = book.tables.get(name)
        if not table or table.blocked:
            return []
        return list(table.rows)

    book.by_logical["表A"] = take(TABLE_A_SHEET)
    book.by_logical["表B"] = take("表 B 成本侧活动扩展表")
    book.by_logical["补登"] = take("流程补登清单")
    book.by_logical["功能"] = take("功能登记表")
    book.by_logical["机能"] = take("机能登记表")
    book.by_logical["不计入"] = take("不计入投入清单")
    book.by_logical["明细"] = take("功能点计算表·数据移动明细")
    book.by_logical["汇总"] = take("功能点计算表·机能规模汇总")
    book.by_logical["链路"] = take("链路核对视图")
    book.by_logical["登记簿"] = take("能力登记簿")
    book.by_logical["清单"] = take("全局活动清单")

    alloc = []
    for name, table in book.tables.items():
        if table.blocked:
            continue
        if name == "分摊表" or name.startswith("分摊表·"):
            alloc.extend(table.rows)
    book.by_logical["分摊"] = alloc

    identity = []
    system = []
    for name, table in book.tables.items():
        if table.blocked:
            continue
        if name == "全系统成本恒等式":
            for row in table.rows:
                row.case = "全系统"
            system.extend(table.rows)
        elif name == "成本恒等式" or name.startswith("成本恒等式·"):
            for row in table.rows:
                if row.case is None and row.get("范围") not in ("", "本片合计"):
                    row.case = row.get("范围")
            identity.extend(table.rows)
    book.by_logical["恒等式"] = identity
    book.by_logical["全系统"] = system
    copies = take("已结期次留存副本")
    book.by_logical["副本"] = copies


def load_previous(path: Path | None, book: WorkbookData) -> None:
    if path is None:
        book.previous = None
        return
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers, raw, _missing, role_cols, group_found = read_table_a_grid(ws, require_separator=False)
    book.previous = _records(ws.title, headers, raw)
    book.previous_headers = headers
    book.previous_role_cols = role_cols
    book.previous_group_found = group_found
    wb.close()


def load_copy(path: Path | None, book: WorkbookData) -> None:
    if path is None:
        if book.by_logical.get("副本"):
            book.copy_rows = book.by_logical["副本"]
        else:
            book.copy_rows = None
        return
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers, raw, _missing = read_grid(ws, skip_note=True, require_separator=False)
    book.copy_rows = _records(ws.title, headers, raw)
    wb.close()


def load_catalog(directory: Path | None, book: WorkbookData) -> None:
    if directory is None:
        book.catalog = None
        return
    found = {}
    if directory.is_dir():
        for item in directory.iterdir():
            if item.stem in EXPORT_FILES and item.suffix.lower() in {".txt", ".csv", ".md", ""}:
                found[item.stem] = item
    book.catalog_missing_files = [name for name in EXPORT_FILES if name not in found]
    keys = set()
    for path in found.values():
        for raw_line in path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line in {"名录项", "主键"} or line.startswith("#"):
                continue
            keys.add(unkey(line))
    book.catalog = keys


def blocked_names(book: WorkbookData) -> set[str]:
    blocked = set()
    for name, table in book.tables.items():
        if table.blocked:
            blocked.add(name)
    return blocked
