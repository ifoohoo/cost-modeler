# -*- coding: utf-8 -*-
"""语义候选。只筛行，不下判不过。"""
from __future__ import annotations

import re

from openpyxl.utils import get_column_letter

from .load import WorkbookData
from .model import LOGICAL_SHEET
from .textutil import callers, fset, query_spans, unkey

ENGLISH = re.compile(r"[A-Za-z][A-Za-z0-9_]{1,}")
PATH = re.compile(r"/[A-Za-z0-9_\-./]+")


def _add(items, sheet, row, column, content, kind):
    items.append({"表名": sheet, "行号": row, "列名": column, "内容": content, "疑点类别": kind})


def _business(items, sheet, row, column, content):
    if not content or content == "—":
        return
    if ENGLISH.search(content) or PATH.search(content):
        _add(items, sheet, row, column, content, "业务语义红线")


def _affirmative_labels(span: str) -> int:
    """查询结果里完整的肯定标签。子串、否定和说不准的说法不算。"""
    text = span.replace("`", "")
    found = 0
    for matched in re.finditer(r"沿用通用字典|沿用通用参数", text):
        previous = text[matched.start() - 1] if matched.start() else ""
        following = text[matched.end():matched.end() + 1]
        if previous not in ("", "为", "，", "。", "：", "；"):
            continue
        if following not in ("", "；", "，", "。", "："):
            continue
        if previous == "为":
            before = text[max(0, matched.start() - 3):matched.start() - 1]
            if before.endswith("不") or before.endswith("是否"):
                continue
        found += 1
    return found


def collect(book: WorkbookData) -> list[dict]:
    items = []
    for rec in book.logical("表A"):
        for column in ("活动名称", "输入", "输出"):
            _business(items, rec.sheet, rec.row, column, rec.get(column, ""))
    for rec in book.logical("表B"):
        for column in ("输出：业务对象 + 完成后的状态", "输入：业务对象 + 进入时的状态"):
            _business(items, rec.sheet, rec.row, column, rec.get(column, ""))
    for rec in book.logical("补登"):
        _business(items, rec.sheet, rec.row, "候选活动名称", rec.get("候选活动名称", ""))
    for rec in book.logical("机能"):
        matched = re.search(r"区块业务目的：([^；]*)", rec.get("来源证据", ""))
        if not matched:
            continue
        purpose = matched.group(1).strip()
        _business(items, rec.sheet, rec.row, "来源证据", purpose)
        if purpose in ("", "—") or len(purpose) < 4:
            _add(items, rec.sheet, rec.row, "来源证据", purpose, "证据叙述不清")
    for rec in book.logical("功能"):
        evidence = rec.get("来源证据", "")
        matched = re.search(r"非业务功能理由：([^；]*)", evidence)
        if matched and len(matched.group(1).strip()) < 8:
            _add(items, rec.sheet, rec.row, "来源证据", matched.group(1).strip(), "证据叙述不清")
        for span in query_spans(evidence):
            if span.count("沿用") > _affirmative_labels(span):
                _add(items, rec.sheet, rec.row, "来源证据", evidence, "沿用措辞含糊")
                break
    for item in items:
        table = book.tables[item["表名"]]
        column = table.headers.index(item["列名"]) + 1
        item["单元格"] = f"{get_column_letter(column)}{item['行号']}"
        rec = next(rec for rec in table.rows if rec.row == item["行号"])
        item["单元格原文"] = rec.get(item["列名"], "")
    return items


def _ref(rec) -> dict:
    return {"表名": rec.sheet, "行号": rec.row}


def _row_context(book: WorkbookData, rec) -> dict:
    headers = book.tables[rec.sheet].headers
    return {
        **_ref(rec),
        "对象标识": rec.get("编码", rec.get("主键", rec.get("机能主键", ""))),
        "单元格": [
            {"列名": header, "坐标": f"{get_column_letter(column)}{rec.row}", "原文": rec.get(header, "")}
            for column, header in enumerate(headers, start=1)
        ],
    }


def _table_context(book: WorkbookData, logical: str) -> dict:
    name = LOGICAL_SHEET[logical]
    names = [name]
    if logical in ("分摊", "恒等式"):
        names += [item for item in book.tables if item.startswith(name + "·")]
    sheets = []
    for sheet_name in names:
        table = book.tables.get(sheet_name)
        reasons = []
        if table is None or not table.present:
            reasons.append("工作表缺失")
        elif table.blocked:
            if table.separator_missing:
                reasons.append("填写区与示例区分隔行缺失")
            if table.header_bad:
                reasons.append("表头不合格")
        rows = table.rows if table and not table.blocked else []
        sheets.append({
            "表名": sheet_name,
            "读取状态": "结构阻断" if reasons else ("空登记表" if not rows else "已读取"),
            "阻断原因": reasons,
            "表头": [
                {"坐标": f"{get_column_letter(column)}{3 if logical == '表A' else 2}", "原文": header}
                for column, header in enumerate(table.headers if table and table.present else [], start=1)
            ],
            "行": [_row_context(book, rec) for rec in rows],
        })
    rows = [row for sheet in sheets for row in sheet["行"]]
    reasons = [f"{sheet['表名']}：{reason}" for sheet in sheets for reason in sheet["阻断原因"]]
    return {
        "表名": name,
        "读取状态": "结构阻断" if reasons else ("空登记表" if not rows else "已读取"),
        "阻断原因": reasons,
        "表头": sheets[0]["表头"],
        "行": rows,
        "工作表": [{key: value for key, value in sheet.items() if key != "行"} for sheet in sheets],
    }


def _association_state(value: str) -> str:
    if value == "—":
        return "尚未建立"
    return "未填写" if not value else "已填写"


def _screen_host(key: str) -> str | None:
    parts = unkey(key).split("|")
    return parts[1].split(" ", 1)[0] if len(parts) == 3 and parts[0] == "画面" else None


def association_context(book: WorkbookData) -> dict:
    """完整关联材料；按填写的身份连接行，不替模型判断业务合理性。"""
    activities = book.logical("表B")
    functions = book.logical("功能")
    mechanisms = book.logical("机能")
    details = book.logical("明细")

    def function_refs(code):
        # 重复主键属于确定规则问题，但不能覆盖其中一行的原始正文。
        return [_ref(rec) for rec in functions if rec.get("主键") == code]

    activity_links = []
    for rec in activities:
        value = rec.get("关联业务功能编号（F-）", "")
        activity_links.append({
            "活动": _ref(rec),
            "表A活动": [_ref(row) for row in book.logical("表A") if row.get("编码") == rec.get("主键")],
            "关联单元格": f"G{rec.row}",
            "关联原文": value,
            "关联状态": _association_state(value),
            "功能": [{"编号": code, "对应行": function_refs(code)} for code in fset(value)],
        })

    mechanism_links = []
    for rec in mechanisms:
        key = unkey(rec.get("主键", ""))
        owner = rec.get("主归属功能编号", "")
        direct_details = [row for row in details if unkey(row.get("机能主键", "")) == key]
        host = _screen_host(key)
        shared_details = [row for row in details if host and row.get("共享加载画面标识") == host]
        trigger_codes = sorted({code for row in direct_details for code in fset(row.get("触发业务功能编号"))})
        all_codes = set(fset(owner)) | set(trigger_codes)
        beneficiary = re.search(r"受益活动：(.*?)；来源档：", rec.get("受益方留痕", ""))
        beneficiary_codes = fset(beneficiary.group(1)) if beneficiary else []
        caller_keys = callers(rec.get("来源证据", ""))
        mechanism_links.append({
            "机能": _ref(rec),
            "主归属单元格": f"F{rec.row}",
            "主归属原文": owner,
            "主归属状态": _association_state(owner),
            "主归属功能": [{"编号": code, "对应行": function_refs(code)} for code in fset(owner)],
            "数据移动": [_ref(row) for row in direct_details],
            "共享加载上下文": [_ref(row) for row in shared_details],
            "触发功能": [{"编号": code, "对应行": function_refs(code)} for code in trigger_codes],
            "关联活动": [_ref(row) for row in activities if all_codes & set(fset(row.get("关联业务功能编号（F-）")))],
            "受益活动": [
                {"编码": code, "对应行": [_ref(row) for row in activities if row.get("主键") == code]}
                for code in beneficiary_codes
            ],
            "随调用方机能": [
                {"主键": caller, "对应行": [_ref(row) for row in mechanisms if unkey(row.get("主键")) == caller]}
                for caller in caller_keys
            ],
        })
    return {
        "语义复核状态": "未运行",
        "方法论": ["docs/方法论/05/article.md", "docs/方法论/06/article.md", "docs/教学文章/附录D-填写模板-v3.11.md"],
        "表": {logical: _table_context(book, logical) for logical in LOGICAL_SHEET if logical not in ("清单", "全系统", "副本") or LOGICAL_SHEET[logical] in book.tables},
        "活动功能关联": activity_links,
        "功能机能关联": mechanism_links,
    }
