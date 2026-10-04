# -*- coding: utf-8 -*-
"""填写引导。

--dump 只读指定表的指定行：列名、该列填写说明、格子里已有的字。
不决定问什么，也不打开别的表。
--write 把已经推出的列文字写入该行还空着的格子。不推导，不覆盖已填内容。
表 A、说明、整页级份额（机器执行）、链路核对视图不写。
--facts 不打开工作簿，只按现场事实推出表 B。
--apply 只接受现场事实，只写由此推出的空表 B 格子。
检索式填写用 --dump 和 --write，不用这两条来决定问什么。
不下检查结论。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from checks.sections import load_sections

ROOT = Path(__file__).resolve().parent
PACKAGE_ROOT = ROOT.parent
METHOD = PACKAGE_ROOT / "references" / "填写方法论.md"
SAYINGS = PACKAGE_ROOT / "references" / "引导说法.md"

_NONE = {"没有", "无", "没", "—", "-", "否", "没有别的", "还没挂", "没挂", "还没有"}
_SPLIT = re.compile(r"[；;、,，]")


def _sayings() -> dict[str, str]:
    text = load_sections(SAYINGS)
    missing = [key for key in _SAYING_KEYS if not text.get(key)]
    if missing:
        raise KeyError("引导说法缺少：" + "、".join(missing))
    return text


_SAYING_KEYS = (
    "question.activity_code",
    "question.outputs",
    "question.inputs",
    "question.owner_l4",
    "question.extra_l4",
    "question.where",
    "question.where_again",
    "question.systems",
    "question.multi_wait",
    "question.function_ids",
    "question.function_ids_again",
    "question.functions_drilled",
    "question.evidence_kind",
    "question.evidence_kind_again",
    "question.evidence_name",
    "question.benefit",
    "question.benefit_again",
    "question.benefit_detail",
    "because.主键",
    "because.输出",
    "because.输入",
    "because.主归属",
    "because.主归属引用",
    "because.执行方式.线下",
    "because.执行方式.线上",
    "because.执行方式.并存",
    "because.承载系统.无",
    "because.承载系统.有",
    "because.功能.无",
    "because.功能.有",
    "because.下钻.不适用",
    "because.下钻.已下钻",
    "because.下钻.待下钻",
    "because.来源",
    "because.留痕.不适用",
    "because.留痕.等分",
    "because.留痕.权重",
    "because.设计告警",
    "stop.要等拆开",
)


def heading(prefix: str) -> str:
    """按标题编号取出方法论里的那一行标题。"""
    for line in METHOD.read_text(encoding="utf-8").splitlines():
        if line.startswith(prefix) and (len(line) == len(prefix) or line[len(prefix)] in " \t"):
            return line
    raise KeyError(prefix)


def _ask(say: dict[str, str], key: str, *, again: bool = False) -> dict:
    name = f"question.{key}_again" if again else f"question.{key}"
    if name not in say:
        name = f"question.{key}"
    return {"kind": "ask", "fact_key": key, "question": say[name]}


def _text(value: object) -> str | None:
    if not isinstance(value, str):
        return None
    text = value.strip()
    return text or None


def _parts(text: str) -> list[str]:
    return [part.strip() for part in _SPLIT.split(text) if part.strip()]


def _object_states(value: object) -> str | None:
    text = _text(value)
    if text is None:
        return None
    chunks = [part.strip() for part in re.split(r"[；;]", text) if part.strip()]
    done = []
    for part in chunks:
        if "·" in part:
            obj, _, state = part.partition("·")
            if not obj.strip() or not state.strip() or "·" in state:
                return None
            done.append(f"{obj.strip()}·{state.strip()}")
            continue
        matched = False
        for sep in ("，", ","):
            if sep in part:
                obj, state = part.split(sep, 1)
                if obj.strip() and state.strip():
                    done.append(f"{obj.strip()}·{state.strip()}")
                    matched = True
                break
        if not matched:
            return None
    return "；".join(done) if done else None


def _code(value: object) -> str | None:
    text = _text(value)
    if text is None or text in _NONE:
        return None
    return text


def _extra(value: object) -> list[str] | None:
    text = _text(value)
    if text is None:
        return None
    if text in _NONE:
        return []
    parts = _parts(text)
    if not parts or any(part in _NONE for part in parts):
        return None
    return parts


def _where(value: object) -> str | None:
    text = _text(value)
    if text is None:
        return None
    both = ("两边都", "系统里和系统外", "又在系统", "既在系统", "并存")
    outside = ("系统外", "系统外面", "不在系统", "没有系统", "线下")
    inside = ("系统里", "业务系统", "在系统中", "线上")
    if any(word in text for word in both):
        return "并存"
    outside_hit = any(word in text for word in outside)
    inside_hit = any(word in text for word in inside)
    if outside_hit and inside_hit:
        return "并存"
    if outside_hit:
        return "线下"
    if inside_hit:
        return "线上"
    return None


def _systems(value: object) -> list[str] | None:
    text = _text(value)
    if text is None:
        return None
    parts = _parts(text)
    if not parts or any(part in _NONE for part in parts):
        return None
    return parts


def _wait(value: object) -> bool | None:
    text = _text(value)
    if text is None:
        return None
    if any(word in text for word in ("要等", "等另一个", "等结果")):
        return True
    if any(word in text for word in ("不用等", "各自", "不等")):
        return False
    return None


def _functions(value: object) -> list[str] | None:
    text = _text(value)
    if text is None:
        return None
    if text in _NONE:
        return []
    parts = _parts(text)
    if parts and all(part.startswith("F-") and " " not in part for part in parts):
        return parts
    return None


def _drilled(value: object) -> bool | None:
    text = _text(value)
    if text is None:
        return None
    if any(word in text for word in ("还没", "没拆", "没有", "不都")):
        return False
    if any(word in text for word in ("都拆", "已经拆", "拆完", "都有", "都已经")):
        return True
    return None


def _evidence_kind(value: object) -> str | None:
    text = _text(value)
    if text is None:
        return None
    machine = any(word in text for word in ("导出", "机器", "系统自动", "自动得出"))
    human = any(word in text for word in ("认定", "人看", "人工", "看过"))
    if machine and human:
        return None
    if machine:
        return "机器导出"
    if human:
        return "人工认定"
    return None


def _benefit(value: object) -> str | None:
    text = _text(value)
    if text is None:
        return None
    weighted = any(word in text for word in ("各占", "定过", "指定了", "明确约定", "具体比例", "各自占比", "占比"))
    equal = any(word in text for word in ("平均", "平分", "等分"))
    if weighted and equal:
        return None
    if weighted:
        return "权重"
    if equal:
        return "等分"
    return None


def _weight(value: object) -> str | None:
    text = _text(value)
    if text is None:
        return None
    parts = [part.strip() for part in re.split(r"[；;]", text) if part.strip()]
    if len(parts) != 3:
        return None
    return f"权重：{parts[0]}；理由：{parts[1]}；指定方：{parts[2]}"


def _cell(column: str, value: str, notes: list[str], prefixes: list[str]) -> dict:
    return {
        "column": column,
        "value": value,
        "notes": notes,
        "headings": [heading(prefix) for prefix in prefixes],
    }


def guide(facts: dict | None) -> dict:
    """按已有事实返回下一问，或表 B 这一行的填法。"""
    got = facts or {}
    say = _sayings()

    code = _code(got.get("activity_code"))
    if code is None:
        return _ask(say, "activity_code", again="activity_code" in got)
    outputs = _object_states(got.get("outputs"))
    if outputs is None:
        return _ask(say, "outputs", again="outputs" in got)
    inputs = _object_states(got.get("inputs"))
    if inputs is None:
        return _ask(say, "inputs", again="inputs" in got)
    owner = _code(got.get("owner_l4"))
    if owner is None:
        return _ask(say, "owner_l4", again="owner_l4" in got)
    extra = _extra(got.get("extra_l4")) if "extra_l4" in got else None
    if extra is None:
        return _ask(say, "extra_l4", again="extra_l4" in got)
    where = _where(got.get("where"))
    if where is None:
        return _ask(say, "where", again="where" in got)

    systems: list[str] = []
    waiting = False
    if where in {"线上", "并存"}:
        systems_value = _systems(got.get("systems")) if "systems" in got else None
        if systems_value is None:
            return _ask(say, "systems", again="systems" in got)
        systems = systems_value
        if len(systems) >= 2:
            waiting_value = _wait(got.get("multi_wait")) if "multi_wait" in got else None
            if waiting_value is None:
                return _ask(say, "multi_wait", again="multi_wait" in got)
            waiting = waiting_value
            if waiting:
                return {
                    "kind": "stop",
                    "heading": heading("### 1.2"),
                    "message": say["stop.要等拆开"],
                }

    functions = _functions(got.get("function_ids")) if "function_ids" in got else None
    if functions is None:
        return _ask(say, "function_ids", again="function_ids" in got)
    drilled = False
    if functions:
        drilled_value = _drilled(got.get("functions_drilled")) if "functions_drilled" in got else None
        if drilled_value is None:
            return _ask(say, "functions_drilled", again="functions_drilled" in got)
        drilled = drilled_value

    kind = _evidence_kind(got.get("evidence_kind"))
    if kind is None:
        return _ask(say, "evidence_kind", again="evidence_kind" in got)
    evidence_name = _text(got.get("evidence_name")) if "evidence_name" in got else None
    if evidence_name is None or evidence_name in _NONE:
        return _ask(say, "evidence_name", again="evidence_name" in got)

    if where == "线下" and not functions:
        drill = "不适用"
        benefit = "—"
    else:
        drill = "已下钻" if functions and drilled else "待下钻"
        benefit_kind = _benefit(got.get("benefit")) if "benefit" in got else None
        if benefit_kind is None:
            return _ask(say, "benefit", again="benefit" in got)
        if benefit_kind == "等分":
            benefit = "等分"
        else:
            benefit = _weight(got.get("benefit_detail")) if "benefit_detail" in got else None
            if benefit is None:
                return _ask(say, "benefit_detail", again="benefit_detail" in got)

    if extra:
        belong = f"主归属 L4：{owner}；引用 L4：{'、'.join(extra)}"
        belong_note = say["because.主归属引用"]
    else:
        belong = f"主归属 L4：{owner}"
        belong_note = say["because.主归属"]
    evidence = f"{kind}：{evidence_name}"
    evidence_notes = [say["because.来源"]]
    if len(systems) >= 2 and not waiting:
        evidence += f"；设计告警：{'；'.join(systems)}"
        evidence_notes.append(say["because.设计告警"])

    cells = [
        _cell("主键", code, [say["because.主键"]], ["### 1.2"]),
        _cell("输出：业务对象 + 完成后的状态", outputs, [say["because.输出"]], ["### 0.4", "### 1.2"]),
        _cell("输入：业务对象 + 进入时的状态", inputs, [say["because.输入"]], ["### 0.4", "### 1.2"]),
        _cell("主归属 / 引用", belong, [belong_note], ["### 1.2"]),
        _cell("执行方式", where, [say[f"because.执行方式.{where}"]], ["### 1.2"]),
        _cell(
            "承载系统",
            "—" if where == "线下" else "；".join(systems),
            [say["because.承载系统.无" if where == "线下" else "because.承载系统.有"]],
            ["### 1.2"],
        ),
        _cell(
            "关联业务功能编号（F-）",
            "—" if not functions else "；".join(functions),
            [say["because.功能.无" if not functions else "because.功能.有"]],
            ["### 1.2"],
        ),
        _cell("下钻状态", drill, [say[f"because.下钻.{drill}"]], ["### 1.2"]),
        _cell("来源证据", evidence, evidence_notes, ["### 1.2"]),
        _cell(
            "受益方留痕",
            benefit,
            [say["because.留痕.不适用" if benefit == "—" else "because.留痕.等分" if benefit == "等分" else "because.留痕.权重"]],
            ["### 1.2"],
        ),
    ]
    return {"kind": "fill", "sheet": "表B", "cells": cells}


def run_facts(facts: dict) -> dict:
    """按给出的事实逐项喂入，记下实际问过的每一问。"""
    collected: dict = {}
    asked = []
    while True:
        result = guide(collected)
        if result["kind"] != "ask":
            result["asked"] = asked
            return result
        key = result["fact_key"]
        if key in collected:
            result["asked"] = asked
            result["rejected"] = True
            return result
        asked.append({"fact_key": key, "question": result["question"]})
        if key not in facts:
            result["asked"] = asked
            return result
        collected[key] = facts[key]


GUIDE = PACKAGE_ROOT / "references" / "列填写指引.md"
TABLE_A_SHEET = "表 A 活动明细表"
SHEET_B_SHEET = "表 B 成本侧活动扩展表"
READONLY_SHEETS = {"说明", "整页级份额（机器执行）", TABLE_A_SHEET, "链路核对视图"}
_B_ASK = {
    "主键": "activity_code",
    "输出：业务对象 + 完成后的状态": "outputs",
    "输入：业务对象 + 进入时的状态": "inputs",
    "主归属 / 引用": "owner_l4",
    "执行方式": "where",
    "关联业务功能编号（F-）": "function_ids",
    "来源证据": "evidence_kind",
    "受益方留痕": "benefit",
}
_B_DERIVED = {"承载系统", "下钻状态"}
_FACT_KEYS = {
    "activity_code",
    "outputs",
    "inputs",
    "owner_l4",
    "extra_l4",
    "where",
    "systems",
    "multi_wait",
    "function_ids",
    "functions_drilled",
    "evidence_kind",
    "evidence_name",
    "benefit",
    "benefit_detail",
}
_FIELD_LABELS = ("用途", "谁填", "怎么写", "不手填", "要问")


def load_column_guide(path: Path | None = None) -> dict[str, dict]:
    """按表名读列填写指引。页和表都收进来，列块保留用途、谁填、写法。"""
    sheets: dict[str, dict] = {}
    current: str | None = None
    column: str | None = None
    bucket: list[str] = []

    def flush() -> None:
        nonlocal bucket
        if current is None:
            bucket = []
            return
        target = sheets[current]
        if column is None:
            target["intro"] = "\n".join(bucket).strip()
        else:
            block = target["columns"][column]
            for line in bucket:
                for label in _FIELD_LABELS:
                    prefix = label + "："
                    if line.startswith(prefix):
                        block[label] = line[len(prefix):].strip()
        bucket = []

    for line in (path or GUIDE).read_text(encoding="utf-8").splitlines():
        if line.startswith("## 页：") or line.startswith("## 表："):
            flush()
            name = line.split("：", 1)[1].strip()
            sheets[name] = {
                "kind": "page" if line.startswith("## 页：") else "table",
                "intro": "",
                "columns": {},
            }
            current = name
            column = None
            bucket = []
            continue
        if line.startswith("### "):
            flush()
            column = line[4:].strip()
            if current is None:
                raise KeyError("列块出现在表名之前")
            sheets[current]["columns"][column] = {}
            bucket = []
            continue
        if current is not None:
            bucket.append(line)
    flush()
    return sheets


def _header_pairs(ws, sheet: str) -> list[tuple[int, str]]:
    row = 3 if sheet == TABLE_A_SHEET else 2
    found = []
    for col in range(1, (ws.max_column or 1) + 1):
        value = ws.cell(row, col).value
        text = "" if value is None else str(value).strip()
        if text:
            found.append((col, text))
    return found


def _explain(block: dict) -> str:
    return block.get("怎么写") or block.get("不手填") or block.get("用途") or ""


def dump_row(workbook: Path, sheet: str, row: int) -> dict:
    """读取指定行。列名和填写说明来自这张表自己的表头，不查看别的表。"""
    from openpyxl import load_workbook

    if row < 1:
        return {"kind": "error", "message": "行号从 1 起"}
    wb = load_workbook(workbook, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return {"kind": "error", "message": "工作簿没有这张表"}
        ws = wb[sheet]
        help_row = 4 if sheet == TABLE_A_SHEET else 3
        columns = []
        for col, name in _header_pairs(ws, sheet):
            value = ws.cell(row, col).value
            shown = "" if value is None else str(value).strip()
            help_value = ws.cell(help_row, col).value
            help_text = "" if help_value is None else str(help_value).strip()
            columns.append({"column": name, "value": shown or None, "help": help_text})
        return {"kind": "dump", "sheet": sheet, "row": row, "columns": columns}
    finally:
        wb.close()


def write_cells(workbook: Path, sheet: str, row: int, cells: dict) -> dict:
    """把已经推出的文字写入该行还空着的格子。不推导，不覆盖。"""
    from openpyxl import load_workbook

    if row < 1:
        return {"kind": "error", "message": "行号从 1 起"}
    if not isinstance(cells, dict):
        return {"kind": "error", "message": "写入内容必须是一个对象"}
    wb = load_workbook(workbook)
    try:
        if sheet not in wb.sheetnames:
            return {"kind": "error", "message": "工作簿没有这张表"}
        if sheet in READONLY_SHEETS:
            return {"kind": "stop", "sheet": sheet, "row": row, "cells": [], "skipped": []}
        help_row = 4 if sheet == TABLE_A_SHEET else 3
        if row <= help_row:
            return {"kind": "error", "message": "这一行是表头或填写说明，不写入"}
        ws = wb[sheet]
        first = ws.cell(row, 1).value
        if isinstance(first, str) and first.startswith("以下为教学案例示例"):
            return {"kind": "error", "message": "这一行是示例分隔，不写入"}
        headers = {name: col for col, name in _header_pairs(ws, sheet)}
        missing = [name for name in cells if name not in headers]
        if missing:
            return {"kind": "error", "message": "没有这些列：" + "、".join(missing)}
        written: list[dict] = []
        skipped: list[dict] = []
        for name, value in cells.items():
            text = "" if value is None else str(value).strip()
            if not text:
                skipped.append({"column": name, "reason": "空值不写"})
                continue
            cell = ws.cell(row, headers[name])
            current = "" if cell.value is None else str(cell.value).strip()
            if current:
                skipped.append({"column": name, "reason": "已有内容不覆盖"})
                continue
            cell.value = text
            written.append({"column": name, "value": text})
        if written:
            wb.save(workbook)
        return {"kind": "fill", "sheet": sheet, "row": row, "cells": written, "skipped": skipped}
    finally:
        wb.close()


def inspect_row(workbook: Path, sheet: str, row: int) -> dict:
    """读取指定行已有内容。还缺事实时只返回下一句。"""
    from openpyxl import load_workbook

    if row < 1:
        return {"kind": "error", "message": "行号从 1 起"}
    sheets = load_column_guide()
    if sheet not in sheets:
        return {"kind": "error", "message": "列填写指引没有这张表"}
    wb = load_workbook(workbook, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return {"kind": "error", "message": "工作簿没有这张表"}
        info = sheets[sheet]
        if info["kind"] != "table" or sheet in READONLY_SHEETS:
            return {
                "kind": "stop",
                "sheet": sheet,
                "row": row,
                "writable": False,
                "explanation": info["intro"],
                "known": [],
                "ask": [],
                "derived": [],
            }
        ws = wb[sheet]
        say = _sayings() if sheet == SHEET_B_SHEET else {}
        known: list[dict] = []
        ask: list[dict] = []
        derived: list[dict] = []
        for col, name in _header_pairs(ws, sheet):
            value = ws.cell(row, col).value
            shown = "" if value is None else str(value).strip()
            block = info["columns"].get(name, {})
            if shown:
                known.append({"column": name, "value": shown})
                continue
            if sheet == SHEET_B_SHEET and name in _B_DERIVED:
                derived.append({"column": name, "explanation": _explain(block)})
                continue
            if sheet == SHEET_B_SHEET and name in _B_ASK:
                key = _B_ASK[name]
                ask.append({
                    "column": name,
                    "fact_key": key,
                    "question": say[f"question.{key}"],
                })
                continue
            # 谁填里出现「机器」只说明有些行由机器带出。有要问就仍要问人。
            if block.get("要问"):
                ask.append({"column": name, "question": block["要问"]})
                continue
            derived.append({"column": name, "explanation": _explain(block)})
        if ask:
            first = ask[0]
            result = {
                "kind": "ask",
                "sheet": sheet,
                "row": row,
                "writable": sheet == SHEET_B_SHEET,
                "column": first["column"],
                "question": first["question"],
                "known": known,
                "derived": derived,
                "ask": [first],
            }
            if "fact_key" in first:
                result["fact_key"] = first["fact_key"]
            return result
        return {
            "kind": "context",
            "sheet": sheet,
            "row": row,
            "writable": sheet == SHEET_B_SHEET,
            "known": known,
            "ask": [],
            "derived": derived,
        }
    finally:
        wb.close()


def apply_row(workbook: Path, sheet: str, row: int, updates: dict) -> dict:
    """只把本入口按现场事实推出的空表 B 格子写上。"""
    from openpyxl import load_workbook

    if row < 1:
        return {"kind": "error", "message": "行号从 1 起"}
    if not isinstance(updates, dict):
        return {"kind": "error", "message": "写入内容必须是一个对象"}
    sheets = load_column_guide()
    if sheet not in sheets:
        return {"kind": "error", "message": "列填写指引没有这张表"}
    info = sheets[sheet]
    if sheet != SHEET_B_SHEET:
        return {
            "kind": "stop",
            "sheet": sheet,
            "row": row,
            "explanation": info["intro"],
            "cells": [],
        }
    foreign = [key for key in updates if key not in _FACT_KEYS]
    if foreign:
        return {
            "kind": "error",
            "message": "只能提交现场事实，不能直接写列：" + "、".join(foreign),
        }
    derived = guide(updates)
    if derived["kind"] != "fill":
        refused = dict(derived)
        refused["sheet"] = sheet
        refused["row"] = row
        refused["cells"] = []
        return refused
    wb = load_workbook(workbook)
    try:
        if sheet not in wb.sheetnames:
            return {"kind": "error", "message": "工作簿没有这张表"}
        ws = wb[sheet]
        first = ws.cell(row, 1).value
        if isinstance(first, str) and first.startswith("以下为教学案例示例"):
            return {"kind": "error", "message": "这一行是示例分隔，不写入"}
        headers = {name: col for col, name in _header_pairs(ws, sheet)}
        missing = [item["column"] for item in derived["cells"] if item["column"] not in headers]
        if missing:
            return {"kind": "error", "message": "没有这些列：" + "、".join(missing)}
        written: list[dict] = []
        skipped: list[dict] = []
        for item in derived["cells"]:
            name = item["column"]
            text = "" if item["value"] is None else str(item["value"]).strip()
            block = info["columns"].get(name, {})
            notes = item.get("notes") or []
            explanation = "".join(notes) if notes else _explain(block)
            if not text:
                skipped.append({"column": name, "reason": "空值不写"})
                continue
            cell = ws.cell(row, headers[name])
            current = "" if cell.value is None else str(cell.value).strip()
            if current:
                skipped.append({"column": name, "reason": "已有内容不覆盖"})
                continue
            cell.value = text
            written.append({"column": name, "value": text, "explanation": explanation})
        if written:
            wb.save(workbook)
        return {"kind": "fill", "sheet": sheet, "row": row, "cells": written, "skipped": skipped}
    finally:
        wb.close()


def _emit(result: dict) -> int:
    json.dump(result, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    if result.get("kind") in {"ask", "error"}:
        return 2
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="按指定的表和行读格子，或把已经推出的空格子写上")
    parser.add_argument("--facts", type=Path, help="JSON 对象，键是事实名，值是现场回答")
    parser.add_argument("--workbook", type=Path, help="要读或写的工作簿")
    parser.add_argument("--sheet", help="表名")
    parser.add_argument("--row", type=int, help="行号，从 1 起")
    parser.add_argument("--dump", action="store_true", help="只读这一行的列名、填写说明和已有的字")
    parser.add_argument("--write", type=Path, help="列名到文字的 JSON。只写还空着的格子")
    parser.add_argument("--apply", type=Path, help="现场事实 JSON。只把由此推出的空表 B 格子写入")
    args = parser.parse_args(argv)
    if args.facts and (args.workbook or args.dump or args.write or args.apply):
        print("事实推导和按行处理不要同时用", file=sys.stderr)
        return 2
    if args.dump and (args.write or args.apply):
        print("读取和写入不要同时用", file=sys.stderr)
        return 2
    if args.write and args.apply:
        print("两种写入不要同时用", file=sys.stderr)
        return 2
    if args.facts:
        facts = json.loads(args.facts.read_text(encoding="utf-8"))
        if not isinstance(facts, dict):
            print("事实文件必须是一个 JSON 对象", file=sys.stderr)
            return 2
        return _emit(run_facts(facts))
    if args.workbook:
        if not args.sheet or not args.row:
            print("按行处理需要表名和行号", file=sys.stderr)
            return 2
        if args.dump:
            return _emit(dump_row(args.workbook, args.sheet, args.row))
        if args.write:
            cells = json.loads(args.write.read_text(encoding="utf-8"))
            return _emit(write_cells(args.workbook, args.sheet, args.row, cells))
        if args.apply:
            updates = json.loads(args.apply.read_text(encoding="utf-8"))
            return _emit(apply_row(args.workbook, args.sheet, args.row, updates))
        return _emit(inspect_row(args.workbook, args.sheet, args.row))
    print("需要事实文件，或工作簿、表名和行号", file=sys.stderr)
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
